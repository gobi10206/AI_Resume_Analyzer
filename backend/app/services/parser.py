import os
import io
import re
from typing import Dict, Any, Tuple
from app.utils.text_cleaning import clean_text, segment_sections

class ResumeParserService:
    @staticmethod
    def parse_file(file_bytes: bytes, filename: str) -> Dict[str, Any]:
        """
        Parse uploaded resume file (PDF, DOCX, TXT) and return extracted text,
        document structure metadata, and section segments.
        """
        ext = os.path.splitext(filename)[1].lower()
        raw_text = ""
        metadata = {
            "file_name": filename,
            "file_type": ext,
            "file_size": len(file_bytes),
            "pages_count": 1,
            "has_tables": False,
            "has_images": False,
            "is_scanned": False,
            "word_count": 0,
            "char_count": 0
        }
        
        if ext == ".pdf":
            raw_text, metadata = ResumeParserService._parse_pdf(file_bytes, metadata)
        elif ext in [".docx", ".doc"]:
            raw_text, metadata = ResumeParserService._parse_docx(file_bytes, metadata)
        elif ext in [".txt", ".rtf", ".md"]:
            raw_text, metadata = ResumeParserService._parse_txt(file_bytes, metadata)
        else:
            raise ValueError(f"Unsupported file format: {ext}. Allowed: PDF, DOCX, TXT.")
            
        cleaned_text = clean_text(raw_text)
        metadata["word_count"] = len(cleaned_text.split())
        metadata["char_count"] = len(cleaned_text)
        
        # Segment into sections
        sections = segment_sections(cleaned_text)
        
        return {
            "raw_text": cleaned_text,
            "metadata": metadata,
            "sections": sections
        }
        
    @staticmethod
    def _parse_pdf(file_bytes: bytes, metadata: Dict[str, Any]) -> Tuple[str, Dict[str, Any]]:
        text_parts = []
        pages_count = 0
        
        # 1. Try pypdf
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            pages_count = len(reader.pages)
            for page in reader.pages:
                extracted = page.extract_text() or ""
                if extracted:
                    text_parts.append(extracted)
        except Exception:
            pass
            
        # 2. If pypdf didn't get enough text, try pdfplumber
        if len("".join(text_parts).strip()) < 50:
            try:
                import pdfplumber
                with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                    pages_count = len(pdf.pages)
                    plumber_texts = []
                    has_tables = False
                    for p in pdf.pages:
                        t = p.extract_text() or ""
                        if t:
                            plumber_texts.append(t)
                        if p.find_tables():
                            has_tables = True
                    if len("".join(plumber_texts).strip()) > len("".join(text_parts).strip()):
                        text_parts = plumber_texts
                        metadata["has_tables"] = has_tables
            except Exception:
                pass
                
        # 3. Check if scanned and try OCR fallback
        combined_text = "\n".join(text_parts).strip()
        if len(combined_text) < 50:
            metadata["is_scanned"] = True
            try:
                import pytesseract
                from PIL import Image
                # If pdf2image is available
                try:
                    from pdf2image import convert_from_bytes
                    images = convert_from_bytes(file_bytes, first_page=1, last_page=3)
                    metadata["has_images"] = True
                    ocr_parts = []
                    for img in images:
                        ocr_text = pytesseract.image_to_string(img)
                        if ocr_text:
                            ocr_parts.append(ocr_text)
                    if ocr_parts:
                        combined_text = "\n".join(ocr_parts)
                except Exception:
                    pass
            except Exception:
                pass
                
        metadata["pages_count"] = max(1, pages_count)
        return combined_text, metadata
        
    @staticmethod
    def _parse_docx(file_bytes: bytes, metadata: Dict[str, Any]) -> Tuple[str, Dict[str, Any]]:
        text_parts = []
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            
            # Paragraphs
            for p in doc.paragraphs:
                if p.text.strip():
                    text_parts.append(p.text.strip())
                    
            # Tables
            if len(doc.tables) > 0:
                metadata["has_tables"] = True
                for table in doc.tables:
                    for row in table.rows:
                        row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_texts:
                            text_parts.append(" | ".join(row_texts))
        except Exception as e:
            text_parts.append(f"[DOCX Parse Fallback] {str(e)}")
            
        return "\n".join(text_parts), metadata

    @staticmethod
    def _parse_txt(file_bytes: bytes, metadata: Dict[str, Any]) -> Tuple[str, Dict[str, Any]]:
        for enc in ["utf-8", "latin-1", "windows-1252", "ascii"]:
            try:
                text = file_bytes.decode(enc)
                return text, metadata
            except UnicodeDecodeError:
                continue
        return file_bytes.decode("utf-8", errors="ignore"), metadata
